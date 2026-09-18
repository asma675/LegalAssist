import shutil
from datetime import datetime
from pathlib import Path
from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
from sqlalchemy import or_
from sqlalchemy.orm import Session
from .config import settings
from .db import Base, engine, get_db
from .models import *
from .schemas import *
from .auth import current_user, verify_password, create_token
from .seed import seed
from .rag import retrieve
from .ai import synthesize
from .intelligence import build_matter_intelligence, intelligence_payload

app = FastAPI(title="Rafi Legal Assist API", version="3.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[x.strip() for x in settings.cors_origins.split(',')],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.on_event("startup")
def startup():
    Path(settings.upload_dir).mkdir(parents=True, exist_ok=True)
    Base.metadata.create_all(bind=engine)
    db = next(get_db())
    try:
        seed(db)
    finally:
        db.close()


@app.get("/api/health")
def health(db: Session = Depends(get_db)):
    return {
        "status": "ok",
        "app": settings.app_name,
        "version": "3.0.0",
        "corpus_documents": db.query(CorpusDocument).count(),
        "demo_mode": settings.demo_mode,
        "matter_intelligence": True,
        "public_early_access": True,
    }


# --- Public early-access endpoints (no authentication required) ---
@app.post("/api/public/waitlist")
def join_waitlist(p: WaitlistIn, db: Session = Depends(get_db)):
    # Invisible honeypot catches basic form bots without blocking real visitors.
    if p.website:
        return {"ok": True, "message": "Thanks — you are on the Rafi early-access list."}
    email = str(p.email).lower().strip()
    existing = db.query(WaitlistLead).filter(WaitlistLead.email == email).first()
    if existing:
        return {"ok": True, "already_joined": True, "message": "You are already on the Rafi early-access list."}
    lead = WaitlistLead(
        email=email, full_name=p.full_name, organization=p.organization, role=p.role,
        team_size=p.team_size, country=p.country, referral=p.referral, source="website_waitlist"
    )
    db.add(lead); db.commit(); db.refresh(lead)
    return {"ok": True, "id": lead.id, "message": "Thanks — you are on the Rafi early-access list."}


@app.post("/api/public/demo-request")
def request_demo(p: DemoRequestIn, db: Session = Depends(get_db)):
    if p.website:
        return {"ok": True, "message": "Thanks — your demo request was received."}
    lead = DemoRequest(
        first_name=p.first_name.strip(), last_name=p.last_name.strip(),
        work_email=str(p.work_email).lower().strip(), organization=p.organization.strip(),
        role=p.role, team_size=p.team_size, practice_area=p.practice_area,
        country=p.country, phone=p.phone, message=p.message, source="website_demo"
    )
    db.add(lead); db.commit(); db.refresh(lead)
    return {"ok": True, "id": lead.id, "message": "Thanks — your demo request was received. We will follow up using the work email you provided."}


@app.get("/api/growth/leads")
def growth_leads(db: Session = Depends(get_db), user: User = Depends(current_user)):
    waitlist = db.query(WaitlistLead).order_by(WaitlistLead.created_at.desc()).limit(250).all()
    demos = db.query(DemoRequest).order_by(DemoRequest.created_at.desc()).limit(250).all()
    return {
        "summary": {"waitlist": db.query(WaitlistLead).count(), "demo_requests": db.query(DemoRequest).count()},
        "waitlist": waitlist,
        "demo_requests": demos,
    }


@app.post("/api/auth/login", response_model=AuthOut)
def login(payload: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == payload.email).first()
    if not user or not verify_password(payload.password, user.password_hash):
        raise HTTPException(401, "Invalid credentials")
    return AuthOut(access_token=create_token(user), user=user)


@app.get("/api/auth/me", response_model=UserOut)
def me(user: User = Depends(current_user)):
    return user


def owned(db, model, user):
    return db.query(model).filter(model.tenant_id == user.tenant_id)


@app.get("/api/dashboard")
def dashboard(db: Session = Depends(get_db), user: User = Depends(current_user)):
    matters = owned(db, Matter, user).all()
    tasks = owned(db, Task, user).all()
    docs = owned(db, Document, user).all()
    upcoming = owned(db, Event, user).filter(Event.start_at >= datetime.utcnow()).order_by(Event.start_at).limit(5).all()
    digests = owned(db, MatterDigest, user).order_by(MatterDigest.created_at.desc()).limit(8).all()
    gaps = owned(db, MissingInformation, user).filter(MissingInformation.status == "open").all()
    contradictions = owned(db, EvidenceLink, user).filter(EvidenceLink.relationship == "contradicts").count()
    graph_claims = owned(db, Claim, user).count()
    matter_titles = {m.id: m.title for m in matters}
    return {
        "metrics": {
            "active_matters": sum(m.status == "active" for m in matters),
            "documents": len(docs),
            "open_tasks": sum(t.status != "done" for t in tasks),
            "corpus": db.query(CorpusDocument).count(),
            "hours_saved_estimate": round(len(docs) * 0.55 + len(tasks) * 0.18 + graph_claims * 0.12, 1),
            "matter_claims": graph_claims,
            "contradictions": contradictions,
            "open_gaps": len(gaps),
        },
        "attention": [
            {"id": m.id, "title": m.title, "stage": m.stage, "risk": m.risk_score, "deadline": m.next_deadline}
            for m in sorted(matters, key=lambda x: x.risk_score, reverse=True)[:4]
        ],
        "upcoming": [{"id": e.id, "title": e.title, "start_at": e.start_at, "location": e.location, "matter_id": e.matter_id} for e in upcoming],
        "digests": [{"id": d.id, "matter_id": d.matter_id, "matter_title": matter_titles.get(d.matter_id, "Matter"), "title": d.title, "summary": d.summary, "change_type": d.change_type, "created_at": d.created_at} for d in digests],
        "usage": {"assistant": 78, "agents": 62, "vault": 71, "research": 55, "contracts": 46, "portal": 38, "matter_intelligence": 86},
    }


@app.get("/api/clients")
def list_clients(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, Client, user).order_by(Client.created_at.desc()).all()


@app.post("/api/clients")
def create_client(p: ClientIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    obj = Client(tenant_id=user.tenant_id, **p.model_dump())
    db.add(obj); db.commit(); db.refresh(obj)
    return obj


@app.get("/api/matters")
def list_matters(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, Matter, user).order_by(Matter.updated_at.desc()).all()


@app.post("/api/matters")
def create_matter(p: MatterIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    obj = Matter(tenant_id=user.tenant_id, **p.model_dump())
    db.add(obj); db.commit(); db.refresh(obj)
    build_matter_intelligence(db, user.tenant_id, obj.id, reason="matter_created")
    return obj


@app.get("/api/matters/{matter_id}")
def get_matter(matter_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    m = owned(db, Matter, user).filter(Matter.id == matter_id).first()
    if not m:
        raise HTTPException(404, "Matter not found")
    return {
        "matter": m,
        "documents": owned(db, Document, user).filter(Document.matter_id == matter_id).all(),
        "tasks": owned(db, Task, user).filter(Task.matter_id == matter_id).all(),
        "notes": owned(db, TeamNote, user).filter(TeamNote.matter_id == matter_id).all(),
        "messages": owned(db, Message, user).filter(Message.matter_id == matter_id).order_by(Message.created_at.desc()).limit(30).all(),
    }


@app.get("/api/matters/{matter_id}/intelligence")
def matter_intelligence(matter_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    payload = intelligence_payload(db, user.tenant_id, matter_id)
    if payload is None:
        raise HTTPException(404, "Matter not found")
    return payload


@app.post("/api/matters/{matter_id}/intelligence/rebuild")
def rebuild_intelligence(matter_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    try:
        build_matter_intelligence(db, user.tenant_id, matter_id, reason="manual_refresh")
    except ValueError:
        raise HTTPException(404, "Matter not found")
    return intelligence_payload(db, user.tenant_id, matter_id)


@app.post("/api/matters/{matter_id}/brief")
def matter_brief(matter_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    intel = intelligence_payload(db, user.tenant_id, matter_id)
    if not intel:
        raise HTTPException(404, "Matter not found")
    sources = []
    for e in intel["evidence"][:10]:
        sources.append({"title": f"{e['document_name']} p. {e['page_number']}", "citation": "matter evidence", "url": None, "text": e["quote"]})
    for issue in intel["issues"]:
        for a in issue["authorities"][:2]:
            sources.append({"title": a["title"], "citation": a["citation"], "url": a["url"], "text": a["relevance_note"]})
    prompt = f"Prepare a concise lawyer-review matter brief for {intel['matter'].title}. Separate verified evidence from inference. Include: key facts, chronology, conflicts, legal issues, missing information, deadlines, and recommended next workflow. Do not invent facts or authorities."
    if not settings.openai_api_key:
        s = intel["summary"]
        brief = (
            f"MATTER BRIEF — {intel['matter'].title}\n\n"
            f"Rafi indexed {s['facts']} source-linked facts, {s['claims']} claims, {s['evidence_links']} evidence links, "
            f"{s['contradictions']} contradiction(s), {s['issues']} legal issue(s), and {s['deadlines']} recorded deadline(s).\n\n"
            "Key review point: use the Claims & Evidence view to inspect each factual proposition against the exact source page. "
            "Review the Issues & Authorities mappings for currency and applicability, then resolve open gaps before final advice or filing.\n\n"
            "Human review required. Connect OPENAI_API_KEY for a fuller source-grounded narrative brief."
        )
    else:
        brief = synthesize(prompt, sources)
    return {"brief": brief, "sources": sources, "human_review_required": True}


@app.get("/api/tasks")
def list_tasks(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, Task, user).order_by(Task.created_at.desc()).all()


@app.post("/api/tasks")
def create_task(p: TaskIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = Task(tenant_id=user.tenant_id, **p.model_dump())
    db.add(o); db.commit(); db.refresh(o)
    if o.matter_id:
        build_matter_intelligence(db, user.tenant_id, o.matter_id, reason="task_added")
    return o


@app.patch("/api/tasks/{task_id}/toggle")
def toggle_task(task_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = owned(db, Task, user).filter(Task.id == task_id).first()
    if not o:
        raise HTTPException(404, "Task not found")
    o.status = "done" if o.status != "done" else "todo"
    db.commit()
    return o


@app.get("/api/events")
def list_events(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, Event, user).order_by(Event.start_at).all()


@app.post("/api/events")
def create_event(p: EventIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = Event(tenant_id=user.tenant_id, **p.model_dump())
    db.add(o); db.commit(); db.refresh(o)
    if o.matter_id:
        build_matter_intelligence(db, user.tenant_id, o.matter_id, reason="event_added")
    return o


@app.get("/api/workflows")
def workflows(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, Workflow, user).all()


@app.post("/api/workflows")
def create_workflow(p: WorkflowIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = Workflow(tenant_id=user.tenant_id, **p.model_dump())
    db.add(o); db.commit(); db.refresh(o)
    return o


@app.post("/api/workflows/{workflow_id}/run")
def run_workflow(workflow_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    w = owned(db, Workflow, user).filter(Workflow.id == workflow_id).first()
    if not w:
        raise HTTPException(404, "Workflow not found")
    w.runs += 1; db.commit()
    return {"status": "queued", "workflow": w.name, "steps": w.steps, "requires_human_approval": True}


@app.get("/api/playbooks")
def list_playbooks(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, FirmPlaybook, user).order_by(FirmPlaybook.created_at.desc()).all()


@app.post("/api/playbooks")
def create_playbook(p: PlaybookIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = FirmPlaybook(tenant_id=user.tenant_id, **p.model_dump())
    db.add(o); db.commit(); db.refresh(o)
    return o


def extract_text_from_upload(path: Path):
    pages = []
    try:
        if path.suffix.lower() == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            for idx, p in enumerate(reader.pages, start=1):
                pages.append({"page_number": idx, "text": (p.extract_text() or "")[:50000]})
        elif path.suffix.lower() == ".docx":
            from docx import Document as WordDocument
            text = "\n".join(p.text for p in WordDocument(str(path)).paragraphs)
            pages = [{"page_number": 1, "text": text[:500000]}]
        elif path.suffix.lower() in {".txt", ".md", ".csv"}:
            pages = [{"page_number": 1, "text": path.read_text(errors="ignore")[:500000]}]
    except Exception:
        pages = []
    full = "\n\n".join(p["text"] for p in pages)[:500000]
    return full or None, pages


@app.get("/api/documents")
def documents(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, Document, user).order_by(Document.created_at.desc()).all()


@app.post("/api/documents/upload")
def upload_document(file: UploadFile = File(...), matter_id: int | None = Form(None), category: str = Form("General"), db: Session = Depends(get_db), user: User = Depends(current_user)):
    if matter_id and not owned(db, Matter, user).filter(Matter.id == matter_id).first():
        raise HTTPException(404, "Matter not found")
    safe = "".join(c for c in file.filename if c.isalnum() or c in "._- ")[:180] or "upload"
    tenant_dir = Path(settings.upload_dir) / str(user.tenant_id)
    tenant_dir.mkdir(parents=True, exist_ok=True)
    dest = tenant_dir / f"{int(datetime.utcnow().timestamp())}_{safe}"
    with dest.open("wb") as f:
        shutil.copyfileobj(file.file, f)
    text_content, pages = extract_text_from_upload(dest)
    obj = Document(tenant_id=user.tenant_id, matter_id=matter_id, name=file.filename, category=category, path=str(dest), text_content=text_content,
                   metadata_json={"content_type": file.content_type, "extracted_chars": len(text_content or ""), "pages": len(pages)})
    db.add(obj); db.flush()
    for p in pages:
        db.add(DocumentPage(tenant_id=user.tenant_id, document_id=obj.id, page_number=p["page_number"], text_content=p["text"]))
    db.commit(); db.refresh(obj)
    if matter_id:
        build_matter_intelligence(db, user.tenant_id, matter_id, reason="document_uploaded")
    return obj


@app.get("/api/documents/{document_id}/page/{page_number}")
def document_page(document_id: int, page_number: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    doc = owned(db, Document, user).filter(Document.id == document_id).first()
    if not doc:
        raise HTTPException(404, "Document not found")
    page = owned(db, DocumentPage, user).filter(DocumentPage.document_id == document_id, DocumentPage.page_number == page_number).first()
    text = page.text_content if page else (doc.text_content or "")
    return {"document_id": doc.id, "name": doc.name, "page_number": page_number, "text": text, "category": doc.category, "matter_id": doc.matter_id}


@app.get("/api/documents/{document_id}/file")
def document_file(document_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    doc = owned(db, Document, user).filter(Document.id == document_id).first()
    if not doc:
        raise HTTPException(404, "Document not found")
    if not doc.path or not Path(doc.path).exists():
        raise HTTPException(404, "Original file is unavailable for this demo record")
    return FileResponse(doc.path, filename=doc.name)


@app.post("/api/notes")
def create_note(p: NoteIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = TeamNote(tenant_id=user.tenant_id, matter_id=p.matter_id, author=user.full_name, body=p.body)
    db.add(o); db.commit(); db.refresh(o)
    return o


@app.post("/api/messages")
def create_message(p: MessageIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = Message(tenant_id=user.tenant_id, matter_id=p.matter_id, sender=user.email, recipient=p.recipient, body=p.body, client_visible=p.client_visible)
    db.add(o); db.commit(); db.refresh(o)
    return o


@app.get("/api/research/history")
def research_history(db: Session = Depends(get_db), user: User = Depends(current_user)):
    return owned(db, ResearchItem, user).order_by(ResearchItem.created_at.desc()).limit(50).all()


def source_payload(d: CorpusDocument, limit: int = 1800):
    return {"id": d.id, "title": d.title, "citation": d.citation, "url": d.source_url, "text": d.text_content[:limit], "source_type": d.source_type, "current_to": d.current_to, "verified": bool((d.metadata_json or {}).get("verified_case_name_and_citation") or (d.metadata_json or {}).get("official_source"))}


@app.post("/api/research")
def research(p: ResearchIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    docs = retrieve(db, p.query, p.jurisdiction, p.top_k)
    sources = [source_payload(d) for d in docs]
    answer = synthesize(p.query, sources)
    item = ResearchItem(tenant_id=user.tenant_id, query=p.query, answer=answer, sources=sources, jurisdiction=p.jurisdiction)
    db.add(item); db.commit()
    return {"answer": answer, "sources": sources, "review_required": True, "disclaimer": "Draft research assistance only. Verify authorities, currency, and applicability before reliance."}


@app.post("/api/ai/chat")
def chat(p: ChatIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    docs = retrieve(db, p.message, p.jurisdiction, 6)
    sources = [source_payload(d, 1600) for d in docs]
    matter_context = ""
    if p.matter_id:
        intel = intelligence_payload(db, user.tenant_id, p.matter_id)
        if intel:
            matter_context = f"\nMatter context: {intel['matter'].title}; {intel['matter'].practice_area}; {intel['matter'].summary or ''}. The matter graph currently has {intel['summary']['claims']} claims, {intel['summary']['contradictions']} contradiction(s), and {intel['summary']['missing']} open information gap(s)."
    return {"answer": synthesize(p.message + matter_context, sources), "sources": sources, "human_review_required": True}


@app.post("/api/contracts/review")
def contract_review(p: ContractReviewIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    playbook = p.playbook or "none supplied"
    if p.playbook_id:
        pb = owned(db, FirmPlaybook, user).filter(FirmPlaybook.id == p.playbook_id).first()
        if pb:
            playbook = f"{pb.name}: " + "; ".join(pb.rules)
    prompt = f"Review this contract excerpt under {p.jurisdiction}. Identify clause type, commercial risks, missing protections, negotiation points, and questions for counsel. Do not state unsupported law. Firm playbook: {playbook}\n\nTEXT:\n{p.text[:16000]}"
    docs = retrieve(db, prompt, p.jurisdiction, 5)
    sources = [source_payload(d, 1000) for d in docs]
    return {"analysis": synthesize(prompt, sources), "sources": sources, "review_required": True}


@app.get("/api/conflicts")
def conflicts(name: str, db: Session = Depends(get_db), user: User = Depends(current_user)):
    n = f"%{name}%"
    clients = owned(db, Client, user).filter(or_(Client.name.ilike(n), Client.company.ilike(n))).all()
    matters = owned(db, Matter, user).filter(or_(Matter.title.ilike(n), Matter.opposing_party.ilike(n), Matter.opposing_counsel.ilike(n))).all()
    return {"query": name, "clients": clients, "matters": matters, "requires_professional_review": True}


@app.get("/api/corpus")
def corpus(q: str | None = None, limit: int = 50, db: Session = Depends(get_db), user: User = Depends(current_user)):
    query = db.query(CorpusDocument)
    if q:
        query = query.filter(or_(CorpusDocument.title.ilike(f"%{q}%"), CorpusDocument.citation.ilike(f"%{q}%"), CorpusDocument.text_content.ilike(f"%{q}%")))
    return query.order_by(CorpusDocument.id.desc()).limit(min(limit, 200)).all()


@app.get("/api/portal/matter/{matter_id}")
def portal_matter(matter_id: int, db: Session = Depends(get_db), user: User = Depends(current_user)):
    m = owned(db, Matter, user).filter(Matter.id == matter_id).first()
    if not m:
        raise HTTPException(404, "Matter not found")
    return {"matter": m, "documents": owned(db, Document, user).filter(Document.matter_id == matter_id).all(), "messages": owned(db, Message, user).filter(Message.matter_id == matter_id, Message.client_visible == True).all()}


@app.post("/api/drafting/generate")
def generate_draft(p: DraftIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    matter_context = ""
    if p.matter_id:
        intel = intelligence_payload(db, user.tenant_id, p.matter_id)
        if intel:
            matter_context = f"Matter: {intel['matter'].title}; practice area: {intel['matter'].practice_area}; summary: {intel['matter'].summary or ''}. There are {intel['summary']['contradictions']} unresolved evidence contradiction(s) and {intel['summary']['missing']} open information gap(s); do not silently resolve them."
    prompt = f"Prepare a lawyer-review draft of a {p.document_type} for {p.jurisdiction}. Use only the facts provided. Do not invent names, dates, legal authorities, remedies or procedural requirements. Mark missing information as [REVIEW / INSERT]. {matter_context}\nFacts: {p.facts}\nAdditional instructions: {p.instructions or 'none'}"
    docs = retrieve(db, prompt, p.jurisdiction, 6)
    sources = [source_payload(d, 1200) for d in docs]
    return {"draft": synthesize(prompt, sources), "sources": sources, "status": "draft", "human_review_required": True}


@app.post("/api/documents/analyze-text")
def analyze_text(p: AnalyzeDocumentIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    prompt = f"Analyze the following text for {p.analysis_type}. Separate document facts from legal inference, identify missing information, key dates, obligations, risks and questions for counsel. Jurisdiction: {p.jurisdiction}. TEXT: {p.text[:18000]}"
    docs = retrieve(db, prompt, p.jurisdiction, 5)
    sources = [source_payload(d, 900) for d in docs]
    return {"analysis": synthesize(prompt, sources), "sources": sources, "human_review_required": True}


@app.post("/api/signatures")
def request_signature(p: SignatureIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    doc = owned(db, Document, user).filter(Document.id == p.document_id).first()
    if not doc:
        raise HTTPException(404, "Document not found")
    req = SignatureRequest(tenant_id=user.tenant_id, document_id=doc.id, signer_email=p.signer_email, status="sent")
    db.add(req); db.commit(); db.refresh(req)
    return req


@app.patch("/api/recommendations/{recommendation_id}")
def update_recommendation(recommendation_id: int, p: RecommendationStatusIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = owned(db, Recommendation, user).filter(Recommendation.id == recommendation_id).first()
    if not o:
        raise HTTPException(404, "Recommendation not found")
    o.status = p.status; db.commit(); db.refresh(o)
    return o


@app.patch("/api/missing-information/{gap_id}")
def update_missing_info(gap_id: int, p: MissingInfoStatusIn, db: Session = Depends(get_db), user: User = Depends(current_user)):
    o = owned(db, MissingInformation, user).filter(MissingInformation.id == gap_id).first()
    if not o:
        raise HTTPException(404, "Gap not found")
    o.status = p.status; db.commit(); db.refresh(o)
    return o


@app.get("/api/templates")
def templates(user: User = Depends(current_user)):
    return [
        {"id": "client-update", "name": "Client status update", "category": "Correspondence", "description": "Matter-grounded client update with approval checklist."},
        {"id": "research-memo", "name": "Research memorandum", "category": "Research", "description": "Issue / rule / analysis / conclusion with source table."},
        {"id": "contract-issues", "name": "Contract issue list", "category": "Contracts", "description": "Clause-by-clause risk and negotiation table."},
        {"id": "chronology", "name": "Evidence chronology", "category": "Litigation", "description": "Dated events linked to exact supporting evidence pages."},
        {"id": "matter-brief", "name": "Matter intelligence brief", "category": "Matter Intelligence", "description": "Facts, conflicts, issues, gaps, deadlines, authorities and next actions."},
        {"id": "demand-letter", "name": "Demand letter draft", "category": "Drafting", "description": "Fact-limited draft with placeholders for lawyer review."},
    ]


@app.get("/api/billing/status")
def billing_status(user: User = Depends(current_user)):
    return {"plan": "professional-demo", "status": "active" if settings.demo_mode else "configure_stripe", "limits": {"research": 500, "documents": 500, "storage_gb": 50}}


@app.post("/api/billing/checkout")
def billing_checkout(user: User = Depends(current_user)):
    if not settings.stripe_secret_key or not settings.stripe_price_professional:
        raise HTTPException(503, "Stripe is not configured. Set STRIPE_SECRET_KEY and STRIPE_PRICE_PROFESSIONAL.")
    import stripe
    stripe.api_key = settings.stripe_secret_key
    session = stripe.checkout.Session.create(
        mode="subscription",
        line_items=[{"price": settings.stripe_price_professional, "quantity": 1}],
        success_url="http://localhost:8080/app?billing=success",
        cancel_url="http://localhost:8080/app?billing=cancelled",
        customer_email=user.email,
        metadata={"tenant_id": str(user.tenant_id), "user_id": str(user.id)},
    )
    return {"url": session.url}


@app.get("/api/calendar/google/connect-url")
def google_connect_url(user: User = Depends(current_user)):
    if not settings.google_client_id or not settings.google_client_secret:
        raise HTTPException(503, "Google Calendar OAuth is not configured.")
    from google_auth_oauthlib.flow import Flow
    flow = Flow.from_client_config(
        {"web": {"client_id": settings.google_client_id, "client_secret": settings.google_client_secret,
                 "auth_uri": "https://accounts.google.com/o/oauth2/auth", "token_uri": "https://oauth2.googleapis.com/token",
                 "redirect_uris": [settings.google_redirect_uri]}},
        scopes=["https://www.googleapis.com/auth/calendar.events"],
    )
    flow.redirect_uri = settings.google_redirect_uri
    url, _ = flow.authorization_url(access_type="offline", include_granted_scopes="true", state=f"{user.id}:{user.tenant_id}")
    return {"url": url, "note": "Persist OAuth tokens encrypted before enabling automatic sync in production."}
